"""Image-backed math/graphs with editable surrounding text in each exporter."""
import io
import re
import subprocess
from xml.sax.saxutils import escape
from visual_assets import Asset, formula, graph

_DELIMITED=re.compile(r'\$\$.*?\$\$|\\\[.*?\\\]|\\\(.*?\\\)|(?<!\\)\$(?!\$).*?(?<!\\)\$',re.S)


def math_asset(source, alt='', display=False):
    try:return formula(source,alt,display)
    except (ValueError,RuntimeError,OSError,TimeoutError,subprocess.TimeoutExpired) as error:return f'[Formula non visualizzabile: {alt or source}]'


def parts(value):
    if isinstance(value,dict):
        if isinstance(value.get('segments'),list):
            result=[]
            for segment in value['segments']:
                if segment.get('type')=='formula':result.append(math_asset(segment.get('latex',''),segment.get('alt_text','')))
                else:result.extend(parts(segment.get('content','')))
            return result
        return parts(value.get('content',''))
    value=str(value or '');result=[];cursor=0
    for match in _DELIMITED.finditer(value):
        if match.start()>cursor:result.append(value[cursor:match.start()])
        raw=match.group();size=2 if raw.startswith(('$$',r'\[',r'\(')) else 1
        result.append(math_asset(raw[size:-size],display=raw.startswith(('$$',r'\['))))
        cursor=match.end()
    if cursor<len(value):result.append(value[cursor:])
    return result or ['']


def events(container, audience):
    for block in container.get('blocks',[]):
        if not isinstance(block,dict):yield 'text',parts(block);continue
        kind=block.get('block_type','text')
        if kind=='formula':yield 'text',[math_asset(block.get('latex',''),block.get('alt_text',''),True)]
        elif kind=='graph':
            try:yield 'text',[graph(block)]
            except (ValueError,KeyError,TypeError,SyntaxError,OverflowError):yield 'text',parts('[Grafico non disponibile] '+str(block.get('alt_text','')))
        elif kind in ('image','image_request'):
            from generated_assets import cached_image
            asset=cached_image(block)
            yield 'text',[asset] if asset else parts('[Immagine non disponibile] '+str(block.get('alt_text',block.get('description',''))))
        elif kind=='list':
            for item in block.get('items',[]):yield 'text',['- ']+parts(item)
        elif kind=='table':
            rows=[block.get('headers',[])]+block.get('rows',[])
            rows=[r if isinstance(r,list) else [r] for r in rows]
            width=max((len(r) for r in rows),default=0)
            if width:yield 'table',[[parts(cell) for cell in row]+[['']]*(width-len(row)) for row in rows]
        elif kind=='speaker_note':yield 'note',parts(block)
        else:yield 'text',parts(block)


def fit(asset,width,scale=1):
    factor=min(scale,width/max(asset.width,1))
    return asset.width*factor,asset.height*factor


def export(artifact,fmt):
    title=str(artifact.get('title','Artefatto'))
    containers=artifact.get('sections') or artifact.get('slides') or []
    if not containers:containers=[{'title':title,'blocks':[{'content':str(artifact.get('content','[Nessun contenuto]'))}]}]
    audience=artifact.get('audience','teacher')
    if fmt=='pdf':data=pdf(title,containers,audience);mime='application/pdf'
    elif fmt=='docx':data=docx(title,containers,audience);mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    elif fmt=='pptx':data=pptx(title,containers,audience);mime='application/vnd.openxmlformats-officedocument.presentationml.presentation'
    else:raise ValueError('Formato non supportato')
    name=re.sub(r'[^\w .-]','_',str(artifact.get('filename_hint','artefatto'))).rsplit('.',1)[0] or 'artefatto'
    return data,mime,name+'.'+fmt


def pdf(title,containers,audience):
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,Image
    from reportlab.lib.styles import getSampleStyleSheet
    styles=getSampleStyleSheet();styles['BodyText'].fontSize=11;styles['BodyText'].leading=16;styles['BodyText'].autoLeading='max'
    output=io.BytesIO();document=SimpleDocTemplate(output,rightMargin=42,leftMargin=42,topMargin=42,bottomMargin=42)
    width=document.width
    def para(units,available):
        if len(units)==1 and isinstance(units[0],Asset):
            asset=units[0];w,h=fit(asset,available-8)
            picture=Image(str(asset.path),width=w,height=h);picture.hAlign='LEFT'
            return picture
        items=[]
        for item in units:
            if isinstance(item,Asset):
                w,h=fit(item,available-8)
                items.append(f'<img src="{escape(str(item.path),{chr(34):"&quot;"})}" width="{w}" height="{h}" valign="middle"/>')
            else:items.append(escape(str(item)).replace('\n','<br/>'))
        from reportlab.lib.styles import ParagraphStyle
        max_height=max([item.height for item in units if isinstance(item,Asset)]+[11])
        style=ParagraphStyle('InlineMath',parent=styles['BodyText'],leading=max(16,max_height+5),spaceBefore=max(0,max_height-11),autoLeading='')
        return Paragraph(''.join(items),style)
    story=[Paragraph(escape(title),styles['Title'])]
    for container in containers:
        story.append(Paragraph(escape(str(container.get('title',''))),styles['Heading2']))
        for kind,content in events(container,audience):
            if kind=='note' and audience!='teacher':continue
            if kind=='table':
                col_width=width/len(content[0]);table=Table([[para(cell,col_width-12) for cell in row] for row in content],colWidths=[col_width]*len(content[0]),repeatRows=1,splitByRow=1,splitInRow=1)
                table.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.4,'#b0b8bb'),('VALIGN',(0,0),(-1,-1),'TOP'),('BACKGROUND',(0,0),(-1,0),'#eef4f3'),('BOTTOMPADDING',(0,0),(-1,-1),8)]));story.append(table)
            else:story.append(para(content,width))
            story.append(Spacer(1,8))
    document.build(story)
    return output.getvalue()


def docx(title,containers,audience, document=None):
    from docx import Document
    from docx.shared import Pt
    from docx.oxml import OxmlElement
    if document is None:
        document=Document();document.styles['Normal'].font.name='Calibri';document.styles['Normal'].font.size=Pt(11)
    document.add_heading(title,0)
    available=(document.sections[0].page_width-document.sections[0].left_margin-document.sections[0].right_margin)/12700
    def fill(paragraph,units,width):
        for item in units:
            if isinstance(item,Asset):
                w,h=fit(item,width-8)
                picture=paragraph.add_run().add_picture(str(item.path),width=Pt(w),height=Pt(h))
                picture._inline.docPr.set('descr',item.alt)
            else:paragraph.add_run(str(item))
    for container in containers:
        document.add_heading(str(container.get('title','')),1)
        for kind,content in events(container,audience):
            if kind=='note' and audience!='teacher':continue
            if kind=='table':
                table=document.add_table(rows=len(content),cols=len(content[0]));table.style='Table Grid'
                for ri,row in enumerate(content):
                    for ci,cell in enumerate(row):fill(table.cell(ri,ci).paragraphs[0],cell,available/len(row))
                marker=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(marker)
            else:fill(document.add_paragraph(),content,available)
    output=io.BytesIO();document.save(output);return output.getvalue()


def line_layout(units,width,font=18):
    from reportlab.pdfbase.pdfmetrics import stringWidth
    lines=[];line=[];x=0;height=font*1.4
    def flush():
        nonlocal line,x,height
        if line:lines.append((line,height))
        line=[];x=0;height=font*1.4
    for unit in units:
        if isinstance(unit,Asset):
            w,h=fit(unit,width,font/12)
            if h>360:
                factor=360/h;w*=factor;h*=factor
            if x+w>width and line:flush()
            line.append((x,unit,w,h));x+=w+2;height=max(height,h+5)
        else:
            for word in re.findall(r'\n|[^\S\n]+|[^\s]+',str(unit)):
                if word=='\n':flush();continue
                w=stringWidth(word,'Helvetica',font)
                if w>width:
                    fragments=[word[i:i+max(1,int(width/font))] for i in range(0,len(word),max(1,int(width/font)))]
                else:fragments=[word]
                for piece in fragments:
                    w=stringWidth(piece,'Helvetica',font)
                    if x+w>width and line:flush()
                    line.append((x,piece,w,font*1.4));x+=w
    flush();return lines


def pptx(title,containers,audience):
    from pptx import Presentation
    from pptx.util import Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    presentation=Presentation();presentation.slide_width=Pt(960);presentation.slide_height=Pt(540)
    def textbox(slide,text,x,y,w,h,size):
        shape=slide.shapes.add_textbox(Pt(x),Pt(y),Pt(max(w+3,4)),Pt(h+4));tf=shape.text_frame
        tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.word_wrap=False
        tf.paragraphs[0].text=text;tf.paragraphs[0].font.size=Pt(size);tf.paragraphs[0].font.name='Arial';return shape
    for container in containers:
        heading=str(container.get('title',title));slide=None;y=0
        def new_slide():
            nonlocal slide,y
            slide=presentation.slides.add_slide(presentation.slide_layouts[6])
            title_shape=textbox(slide,heading,40,20,880,70,30);title_shape.text_frame.word_wrap=True;y=100
        def draw_lines(lines,x,width):
            nonlocal y
            for line,height in lines:
                if y+height>500:new_slide()
                for dx,item,w,h in line:
                    if isinstance(item,Asset):
                        if h>390:
                            factor=390/h;w*=factor;h*=factor
                        picture=slide.shapes.add_picture(str(item.path),Pt(x+dx),Pt(y+(height-h)/2),width=Pt(w),height=Pt(h))
                        picture._element.nvPicPr.cNvPr.set('descr',item.alt)
                    else:textbox(slide,item,x+dx,y+(height-h)/2,w,h,18)
                y+=height
        new_slide()
        for kind,content in events(container,audience):
            if kind=='note':
                tf=slide.notes_slide.notes_text_frame
                tf.text += '\n'+''.join(item.alt if isinstance(item,Asset) else str(item) for item in content)
                continue
            if kind=='table':
                n=len(content[0]);cellwidth=880/n
                for row in content:
                    layouts=[line_layout(cell,cellwidth-16) for cell in row]
                    while any(layouts):
                        if y>420:new_slide()
                        room=500-y-16
                        chunks=[]
                        for layout in layouts:
                            chunk=[];used=0
                            while layout and used+layout[0][1]<=room:
                                line=layout.pop(0);chunk.append(line);used+=line[1]
                            chunks.append(chunk)
                        if not any(chunks):
                            new_slide();continue
                        height=max(sum(h for _,h in chunk) for chunk in chunks)+16
                        start=y
                        for ci,chunk in enumerate(chunks):
                            shape=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(40+ci*cellwidth),Pt(start),Pt(cellwidth),Pt(height))
                            shape.fill.background();shape.line.color.rgb=RGBColor.from_string('B0B8BB')
                            y=start+8;draw_lines(chunk,48+ci*cellwidth,cellwidth-16)
                        y=start+height+4
                        if any(layouts):new_slide()
            else:
                draw_lines(line_layout(content,880),40,880);y+=10
    output=io.BytesIO();presentation.save(output);return output.getvalue()
