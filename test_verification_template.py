import io,unittest,hashlib
from docx import Document
from verification_template import FILE, compile_test

class TemplateTests(unittest.TestCase):
    def test_scores_in_text_blocks(self):
        points=[10,15,20,15,20,10,10]
        blocks=[{'block_type':'text','segments':[
            {'type':'text','content':f'Esercizio {i} ('},
            {'type':'text','content':f'{score} punti): Risolvi.'}]} for i,score in enumerate(points,1)]
        artifact={'title':'Verifica','sections':[{'title':'Verifica di Matematica','blocks':blocks}]}
        content,_,_=compile_test(artifact)
        doc=Document(io.BytesIO(content))
        self.assertNotIn('{{PUNTI',doc._element.xml)
        self.assertIn('Esercizio 7 (10 punti): Risolvi.', '\n'.join(p.text for p in doc.paragraphs))
        blocks.append({'block_type':'text','content':'Esercizio 8: Risolvi.'})
        with self.assertRaisesRegex(ValueError,'Punteggio mancante'):
            compile_test(artifact)

    def test_score_total_reports_actual_value(self):
        with self.assertRaisesRegex(ValueError,'rilevati è 90'):
            compile_test({'sections':[{'title':'Esercizio 1 (90 punti)','blocks':[]}]})

    def test_scores_from_exercise_titles(self):
        artifact={'title':'Prova','sections':[
            {'title':f'Esercizio {i} (Punti: {score})','blocks':[]}
            for i,score in enumerate([10,12,18,15,15,15,15],1)]}
        original=FILE.read_bytes()
        content,_,_=compile_test(artifact)
        xml=Document(io.BytesIO(content))._element.xml
        self.assertNotIn('{{PUNTI',xml)
        self.assertIn('100',xml)
        self.assertEqual(FILE.read_bytes(),original)

    def test_invalid_title_scores_rejected(self):
        for titles in [
            ['Esercizio 1 (Punti: 90)'],
            ['Esercizio 2 (Punti: 100)'],
            ['Esercizio 1 (Punti: 50)','Esercizio 1 (Punti: 50)'],
            ['Esercizio 1 (Punti: 100)','Esercizio 2'],
            ['Esercizio 1 (Punti: 0)','Esercizio 2 (Punti: 100)'],
        ]:
            with self.subTest(titles=titles),self.assertRaises(ValueError):
                compile_test({'sections':[{'title':t,'blocks':[]} for t in titles]})

    def test_fills_and_preserves_original(self):
        original=FILE.read_bytes()
        cell=lambda s:{'segments':[{'type':'text','content':s}]}
        artifact={'title':'Prova','sections':[{'title':'Esercizio 1','blocks':[{'block_type':'text','content':'Risolvi x+1=2.'},{'block_type':'table','headers':[cell('Esercizio'),cell('Punteggio massimo')],'rows':[[cell('1'),cell('100')],[cell('TOTALE'),cell('100')]]}]}]}
        content,_,_=compile_test(artifact)
        self.assertEqual(FILE.read_bytes(),original)
        doc=Document(io.BytesIO(content));xml=doc._element.xml
        self.assertNotIn('{{PUNTI',xml);self.assertNotIn('{{INIZIO',xml)
        self.assertIn('Risolvi x+1=2.',xml);self.assertIn('w:type="page"',xml)
        artifact['sections'][0]['blocks'][1]['rows'][0][1]=cell('90')
        with self.assertRaises(ValueError):compile_test(artifact)
