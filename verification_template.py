"""Validate and preserve a user-supplied DOCX template without modifying it."""
import io
import json
import hashlib
import re
import threading
from pathlib import Path
from zipfile import ZipFile, BadZipFile
from lxml import etree

FILE=Path(__file__).resolve().parent/'verification_template.docx'
LOCK=threading.Lock()
MAX_SIZE=10*1024*1024
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}


def validate(content):
    if not content or len(content)>MAX_SIZE:raise ValueError('File vuoto o superiore a 10 MB.')
    try:
        with ZipFile(io.BytesIO(content)) as z:
            if sum(i.file_size for i in z.infolist())>50*1024*1024:raise ValueError('Documento troppo grande una volta decompresso.')
            root=etree.fromstring(z.read('word/document.xml'),etree.XMLParser(resolve_entities=False,no_network=True))
    except (BadZipFile,KeyError,etree.XMLSyntaxError) as e:raise ValueError('Il file non e un documento DOCX valido.') from e
    paragraphs=[''.join(p.xpath('.//w:t/text()',namespaces=NS)).strip() for p in root.xpath('//w:p',namespaces=NS)]
    normalized=[p.replace('INIZIO ESERCIZI','INIZIO_ESERCIZI').replace('FINE ESERCIZI','FINE_ESERCIZI') for p in paragraphs]
    for marker in ('{{INIZIO_ESERCIZI}}','{{FINE_ESERCIZI}}'):
        if normalized.count(marker)!=1:raise ValueError(marker+' deve comparire una sola volta, in un paragrafo separato.')
    if normalized.index('{{INIZIO_ESERCIZI}}')>=normalized.index('{{FINE_ESERCIZI}}'):raise ValueError('Marcatori degli esercizi in ordine errato.')
    text='\n'.join(paragraphs)
    required=[f'{{{{PUNTI_{i}}}}}' for i in range(1,11)]+['{{PUNTI_TOTALE}}']
    missing=[m for m in required if text.count(m)!=1]
    if missing:raise ValueError('Marcatori punteggi mancanti o duplicati: '+', '.join(missing))
    return {'configured':True,'exercise_slots':10,'size':len(content),'filename':'verification_template.docx'}


def save(content, filename="verification_template.docx"):
    info=validate(content)
    with LOCK:
        temp=FILE.with_suffix('.tmp');temp.write_bytes(content);temp.replace(FILE)
        name=Path(filename.replace('\\','/')).name
        meta={'name':name,'sha256':hashlib.sha256(content).hexdigest()}
        FILE.with_suffix('.json').write_text(json.dumps(meta,ensure_ascii=False),encoding='utf8')
    return status()


def status():
    if not FILE.exists():return {'configured':False}
    content=FILE.read_bytes();info=validate(content)
    info['original_name']=None
    try:
        meta=json.loads(FILE.with_suffix('.json').read_text(encoding='utf8'))
        if meta['sha256']==hashlib.sha256(content).hexdigest():info['original_name']=meta['name']
    except (OSError,ValueError,KeyError):pass
    return info


def compile_test(artifact):
    import copy
    from docx import Document
    from docx.enum.text import WD_BREAK
    from image_renderers import docx
    def text(cell):
        if isinstance(cell,str):return cell
        return ''.join(str(s.get('content',s.get('latex',''))) for s in cell.get('segments',[]))
    containers=copy.deepcopy(artifact.get('sections') or artifact.get('slides') or [])
    scores={};found=False
    for container in containers:
        blocks=[]
        for block in container.get('blocks',[]):
            headers=[text(h).strip().lower() for h in block.get('headers',[])]
            if block.get('block_type')=='table' and 'esercizio' in headers and 'punteggio massimo' in headers:
                if found:raise ValueError('Piu tabelle punteggi: impossibile compilare il modello senza ambiguita.')
                found=True;ni=headers.index('esercizio');pi=headers.index('punteggio massimo')
                for row in block.get('rows',[]):
                    number=text(row[ni]).strip()
                    if number.upper()=='TOTALE':continue
                    match=re.fullmatch(r'(?:Esercizio\s+)?(\d+)',number,re.I)
                    if not match:raise ValueError('Numerazione esercizi non valida nella tabella punteggi.')
                    n=int(match[1]);value=text(row[pi]).strip()
                    if n in scores or not 1<=n<=10 or not value.isdecimal() or int(value)<=0:raise ValueError('Punteggi non validi: massimo 10 esercizi, punti interi positivi.')
                    scores[n]=int(value)
            else:blocks.append(block)
        container['blocks']=blocks
    if not found:
        # Scores can be in section headings or at the start of text paragraphs.
        # Only explicit exercise scores count; never infer or redistribute them.
        headings=[]
        for container in containers:
            headings.append(container.get('title',''))
            for block in container.get('blocks',[]):
                if block.get('block_type')=='text':
                    headings.extend((text(block) or block.get('content','')).splitlines())
        for heading in headings:
            title=heading.strip()
            if not re.match(r'^Esercizio\s+\d+',title,re.I):continue
            match=re.match(r'^Esercizio\s+(\d+)\s*\(\s*(?:Punti\s*:\s*(\d+)|(\d+)\s+punti)\s*\)(?=\s|:|$)',title,re.I)
            if not match:
                raise ValueError('Punteggio mancante o non valido nell\'intestazione: '+title)
            n=int(match[1]);value=int(match[2] or match[3])
            if n in scores or not 1<=n<=10 or value<=0:
                raise ValueError('Punteggi non validi: massimo 10 esercizi, punti interi positivi.')
            scores[n]=value
    if not scores:
        raise ValueError('Nessun punteggio riconosciuto: usare una tabella Esercizio / Punteggio massimo oppure intestazioni come Esercizio 1 (Punti: 10) o Esercizio 1 (10 punti), nei titoli o nel testo. Esportazione standard PDF disponibile.')
    if sorted(scores)!=list(range(1,len(scores)+1)):
        raise ValueError('Numerazione dei punteggi non consecutiva: esercizi rilevati '+', '.join(map(str,sorted(scores)))+'.')
    if sum(scores.values())!=100:
        raise ValueError(f'Il totale dei punteggi rilevati è {sum(scores.values())}, ma il modello Word richiede 100 punti.')
    content=FILE.read_bytes();validate(content);document=Document(io.BytesIO(content));body=document._element.body
    starts=[p for p in document.paragraphs if p.text.strip() in ('{{INIZIO ESERCIZI}}','{{INIZIO_ESERCIZI}}')]
    ends=[p for p in document.paragraphs if p.text.strip() in ('{{FINE ESERCIZI}}','{{FINE_ESERCIZI}}')]
    if len(starts)!=1 or len(ends)!=1:raise ValueError('I marcatori esercizi devono essere paragrafi del corpo del documento, fuori dalle tabelle.')
    start,end=starts[0]._p,ends[0]._p
    originals=set(body)
    docx(artifact.get('title','Verifica'),containers,'students',document=document)
    generated=[el for el in body if el not in originals]
    node=start.getnext()
    while node is not end:
        if node is None:raise ValueError('Intervallo dei marcatori non valido.')
        following=node.getnext();body.remove(node);node=following
    for el in generated:end.addprevious(el)
    body.remove(start)
    # Preserve all material following the marker, starting it on a new page.
    for child in list(end):
        if child.tag!='{'+NS['w']+'}pPr':end.remove(child)
    from docx.text.paragraph import Paragraph
    Paragraph(end,document).add_run().add_break(WD_BREAK.PAGE)
    mapping={f'{{{{PUNTI_{i}}}}}':str(scores[i]) if i in scores else '' for i in range(1,11)}
    mapping['{{PUNTI_TOTALE}}']='100'
    # Markers may span Word runs: replace only their character ranges.
    for paragraph in body.xpath('.//w:p'):
        nodes=paragraph.xpath('.//w:t');joined=''.join(n.text or '' for n in nodes)
        for marker,value in mapping.items():
            pos=joined.find(marker)
            if pos<0:continue
            offset=0
            for n in nodes:
                old=n.text or '';lo=max(0,pos-offset);hi=min(len(old),pos+len(marker)-offset)
                if lo<hi:n.text=old[:lo]+(value if offset<=pos<offset+len(old) else '')+old[hi:]
                offset+=len(old)
            joined=''.join(n.text or '' for n in nodes)
    output=io.BytesIO();document.save(output)
    return output.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document','Verifica-modello.docx'
